"""Reads a processed workbook back into a pandas DataFrame containing only
the *visible* columns (matching what the user would see if they opened the
file in Excel), for display in the Streamlit results tabs.
"""
from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from utils.validators import is_tax_column_header, normalize_header


def read_visible_dataframe(path: Path) -> pd.DataFrame:
    wb = load_workbook(path, data_only=False)
    ws = wb.active

    header_row_idx = None
    headers = None
    for r in range(1, min(ws.max_row, 15) + 1):
        row_vals = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
        if any(is_tax_column_header(v) for v in row_vals):
            header_row_idx = r
            headers = row_vals
            break

    if header_row_idx is None:
        wb.close()
        return pd.DataFrame()

    visible_cols = [
        c for c in range(1, ws.max_column + 1)
        if not ws.column_dimensions[get_column_letter(c)].hidden and headers[c - 1] not in (None, "")
    ]

    records = []
    for r in range(header_row_idx + 1, ws.max_row + 1):
        row_vals = [ws.cell(row=r, column=c).value for c in visible_cols]
        # Skip the Total row (its tax cell is a formula string) and blank rows.
        if any(isinstance(v, str) and v.strip().startswith("=") for v in row_vals):
            continue
        if all(v is None or v == "" for v in row_vals):
            continue
        records.append(row_vals)

    wb.close()

    col_names = [str(headers[c - 1]) for c in visible_cols]
    df = pd.DataFrame(records, columns=col_names)
    return df
