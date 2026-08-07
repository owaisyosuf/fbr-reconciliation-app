"""Sums the 'Sales Tax / FED in ST Mode' column from a processed workbook.

Works directly off raw cell values with openpyxl (not a recalculated
formula) so the total is available immediately in Python, without needing
LibreOffice to open the file first. Formula cells (the Total row itself,
which stores a live =SUM(...) formula) are skipped automatically, as is
anything non-numeric (headers, blanks, text).
"""
from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional

from openpyxl import load_workbook

from utils.exceptions import MissingColumnError, EmptySheetError
from utils.validators import is_tax_column_header
from utils.logging_config import get_logger

logger = get_logger(__name__)


@dataclass
class TotalResult:
    total: float
    row_count: int
    column_header: str
    sheet_name: str
    row_values: list[float] = field(default_factory=list)


def _coerce_number(value) -> Optional[float]:
    """Best-effort conversion of a cell value to float. Returns None if the
    cell isn't a usable number (formula string, text label, blank, etc.).
    """
    if value is None:
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str):
        s = value.strip()
        if not s or s.startswith("="):
            return None
        cleaned = s.replace(",", "")
        try:
            return float(cleaned)
        except ValueError:
            return None
    return None


def calculate_tax_total(path: Path, source_label: str) -> TotalResult:
    """Locate the tax column anywhere in the first 15 rows (handles both the
    FBR layout, header on row 1, and the Local/Annex C layout, header on
    row 3 after Skill 02 processing) and sum every numeric value beneath it.
    """
    wb = load_workbook(path, data_only=False)
    ws = wb.active

    header_row_idx = None
    tax_col_idx = None
    header_text = None

    max_scan_row = min(ws.max_row, 15)
    for r in range(1, max_scan_row + 1):
        for c in range(1, ws.max_column + 1):
            val = ws.cell(row=r, column=c).value
            if is_tax_column_header(val):
                header_row_idx = r
                tax_col_idx = c
                header_text = val
                break
        if header_row_idx:
            break

    if header_row_idx is None or tax_col_idx is None:
        wb.close()
        raise MissingColumnError(
            f"Couldn't find the 'Sales Tax / FED in ST Mode' column in the "
            f"processed {source_label} file."
        )

    values: list[float] = []
    for r in range(header_row_idx + 1, ws.max_row + 1):
        num = _coerce_number(ws.cell(row=r, column=tax_col_idx).value)
        if num is not None:
            values.append(num)

    wb.close()

    if not values:
        raise EmptySheetError(
            f"No numeric 'Sales Tax / FED in ST Mode' values found in the "
            f"{source_label} file after processing."
        )

    total = round(sum(values), 2)
    logger.info(
        "%s total: %.2f across %d rows (column='%s', header_row=%d)",
        source_label, total, len(values), header_text, header_row_idx,
    )
    return TotalResult(
        total=total,
        row_count=len(values),
        column_header=str(header_text),
        sheet_name=ws.title,
        row_values=values,
    )
