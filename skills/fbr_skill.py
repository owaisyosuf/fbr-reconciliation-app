"""Skill 01 -- FBR Processing Skill (wrapper).

Wraps the bundled `format_sales_tax_report.py` logic (Baber Tyre
Corporation's monthly FBR "Domestic Invoices" export formatter) so it can be
called programmatically from the Streamlit app instead of as a standalone
CLI script. The underlying formatting rules are untouched -- see
`skills/fbr_processing/format_sales_tax_report.py`.

What it does to the FBR file:
  1. Never modifies the original upload -- always works on a copy.
  2. Sets font to Calibri throughout.
  3. Hides every column except: Invoice Date, Buyer Name, Quantity, HS Code,
     Value of Sales Excluding Sales Tax, Sales Tax/ FED in ST Mode.
  4. Converts text-formatted numeric columns (e.g. "8,460.00") into real
     numbers.
  5. Adds a Total row with a live =SUM() formula for the tax column.
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from skills.fbr_processing.format_sales_tax_report import format_report, KEEP_HEADERS, SUM_HEADER
from utils.exceptions import MissingColumnError
from utils.logging_config import get_logger

logger = get_logger(__name__)


def apply_fbr_skill(input_path: Path, output_path: Path) -> Path:
    """Apply Skill 01 to the FBR file. Returns the path to the processed copy."""
    logger.info("Applying FBR skill to %s", input_path.name)

    # Sanity-check the target column exists before we bother formatting,
    # so we can raise a clear, user-friendly error instead of a silent no-op.
    wb = load_workbook(input_path, read_only=True)
    ws = wb.active
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    wb.close()

    normalized = {str(h).strip().lower() if h else "" for h in headers}
    if SUM_HEADER.strip().lower() not in normalized:
        raise MissingColumnError(
            "Couldn't find the 'Sales Tax/ FED in ST Mode' column in the FBR file.",
            detail=f"Columns found: {headers}",
        )

    format_report(str(input_path), str(output_path))
    logger.info("FBR skill complete -> %s", output_path.name)
    return output_path
