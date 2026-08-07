"""
Reusable monthly formatter for Baber Tyre Corporation FBR sales invoice reports.
Usage: python format_sales_tax_report.py <input.xlsx> <output.xlsx>

Keeps only: Invoice Date, Buyer Name, Quantity, HS Code,
Value of Sales Excluding Sales Tax, Sales Tax/ FED in ST Mode
(all other columns are hidden, not deleted)
Sets font to Calibri throughout, adds a SUM row for Sales Tax/ FED in ST Mode.
"""
import sys
from openpyxl import load_workbook
from openpyxl.styles import Font

KEEP_HEADERS = {
    "Invoice Date",
    "Buyer Name",
    "Quantity",
    "HS Code",
    "Value of Sales Excluding Sales Tax",
    "Sales Tax/ FED in ST Mode",
}
SUM_HEADER = "Sales Tax/ FED in ST Mode"
# Columns that come in as text like "8,460.00" and need to become real numbers
NUMERIC_HEADERS = {"Quantity", "Value of Sales Excluding Sales Tax", "Sales Tax/ FED in ST Mode"}


def format_report(input_path, output_path):
    wb = load_workbook(input_path)
    ws = wb.active

    max_row = ws.max_row
    max_col = ws.max_column
    header_row = [ws.cell(row=1, column=c).value for c in range(1, max_col + 1)]

    sum_col_idx = None
    for c, header in enumerate(header_row, start=1):
        col_letter = ws.cell(row=1, column=c).column_letter
        if header not in KEEP_HEADERS:
            ws.column_dimensions[col_letter].hidden = True
        if header == SUM_HEADER:
            sum_col_idx = c
        if header in NUMERIC_HEADERS:
            for r in range(2, max_row + 1):
                cell = ws.cell(row=r, column=c)
                if isinstance(cell.value, str):
                    cleaned = cell.value.replace(",", "").strip()
                    if cleaned:
                        try:
                            cell.value = round(float(cleaned), 2)
                            cell.number_format = "#,##0.00"
                        except ValueError:
                            pass

    for row in ws.iter_rows(min_row=1, max_row=max_row, max_col=max_col):
        for cell in row:
            existing = cell.font
            cell.font = Font(
                name="Calibri",
                size=existing.size or 11,
                bold=existing.bold,
                italic=existing.italic,
                color=existing.color,
            )

    if sum_col_idx:
        sum_col_letter = ws.cell(row=1, column=sum_col_idx).column_letter
        total_row = max_row + 1
        label_col_idx = None
        for c, header in enumerate(header_row, start=1):
            if header == "Buyer Name":
                label_col_idx = c
                break
        if label_col_idx:
            label_cell = ws.cell(row=total_row, column=label_col_idx, value="Total")
            label_cell.font = Font(name="Calibri", bold=True)
        sum_cell = ws.cell(
            row=total_row,
            column=sum_col_idx,
            value=f"=SUM({sum_col_letter}2:{sum_col_letter}{max_row})",
        )
        sum_cell.font = Font(name="Calibri", bold=True)

    wb.save(output_path)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print("Usage: python format_sales_tax_report.py <input.xlsx> <output.xlsx>")
        sys.exit(1)
    format_report(sys.argv[1], sys.argv[2])
    print(f"Saved: {sys.argv[2]}")
